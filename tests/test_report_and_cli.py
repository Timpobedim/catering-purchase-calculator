from __future__ import annotations

import re
from decimal import Decimal
from pathlib import Path

import pytest
from conftest import write_book
from openpyxl import load_workbook

from banquet_calc import cli
from banquet_calc.service import run


def purchase_rows(path: Path) -> list[tuple[object, ...]]:
    sheet = load_workbook(path)["Закупка"]
    return [
        row for row in sheet.iter_rows(min_row=5, values_only=True) if any(cell is not None for cell in row)
    ]


def test_demo_banquet_report(demo_files: tuple[Path, Path], tmp_path: Path) -> None:
    menu, catalog = demo_files

    outcome = run(menu, catalog, tmp_path / "out")

    result = outcome.result
    assert outcome.output.name == "Закупка — банкет на 40 гостей.xlsx"
    assert result.margin_percent == Decimal(5)  # из листа «Параметры»
    assert result.warnings == []
    workbook = load_workbook(outcome.output)
    assert workbook.sheetnames == ["Закупка", "По блюдам", "Предупреждения"]

    rows = purchase_rows(outcome.output)
    assert rows[0][:2] == ("1.", "Овощи, фрукты и зелень")
    assert rows[1][0] == "1.1"
    total = rows[-1]
    assert total[1] == "Итого"
    assert total[8] == pytest.approx(float(result.total_cost))

    lines = {line.name: line for line in result.lines}
    # 40 × 0,03 кг соуса × 0,15 кг пармезана + 40 × 0,01 кг — плюс 5 % запаса, упаковка 0,2 кг
    assert lines["Сыр пармезан"].required == Decimal("0.189") + Decimal("0.42")
    assert lines["Сыр пармезан"].packs == 4
    assert lines["Сырная тарелка от партнёра"].packs == 5  # 40 × 0,1 = 4 тарелки, с запасом 4,2 → 5


def test_cli_prints_summary(
    demo_files: tuple[Path, Path], tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    menu, catalog = demo_files

    code = cli.main(
        ["--menu", str(menu), "--catalog", str(catalog), "--out", str(tmp_path), "--margin", "10"]
    )

    out = capsys.readouterr().out
    assert code == 0
    assert "Готово:" in out
    # сумма по-русски: пробел между разрядами, запятая перед копейками
    assert re.search(r"Позиций: 31, сумма закупки: \d{1,3}( \d{3})*,\d{2} ₽", out)


def test_cli_lists_problems_and_exits_with_2(
    demo_files: tuple[Path, Path], tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    _, catalog = demo_files
    menu = write_book(
        tmp_path / "broken.xlsx",
        {
            "Меню": [{"Блюдо": "Борщ", "Порций": 10}],
            "Техкарты": [{"Техкарта": "Щи", "Тип": "блюдо", "Ингредиент": "Капуста", "Количество": 0.1}],
        },
    )

    code = cli.main(["--menu", str(menu), "--catalog", str(catalog), "--out", str(tmp_path)])

    assert code == 2
    assert "блюда «Борщ» нет среди техкарт" in capsys.readouterr().err
