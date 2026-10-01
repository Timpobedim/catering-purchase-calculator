"""Закупочный лист в Excel.

Разделы нумеруются «1.», позиции внутри — «1.1», внизу итог; отдельные листы — расчёт по блюдам
и предупреждения.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from itertools import groupby
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from banquet_calc.calculation import CalculationResult

HEADERS = [
    "№",
    "Наименование",
    "Ед. изм.",
    "Нужно с запасом",
    "Фасовка",
    "Упаковок",
    "К закупке",
    "Цена за уп., ₽",
    "Сумма, ₽",
    "Поставщик",
]
WIDTHS = [7, 42, 9, 16, 10, 11, 12, 14, 14, 26]
QUANTITY_FORMAT = "0.###"
MONEY_FORMAT = "#,##0.00"

BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
SECTION_FILL = PatternFill("solid", fgColor="E8F0E3")
HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _number(value: Decimal | None) -> float | None:
    # Excel хранит числа как float — переводим в самом конце, считали в Decimal.
    return None if value is None else float(value)


def money(value: Decimal) -> str:
    """38460 → «38 460,00» — для окна и командной строки."""
    return f"{value:,.2f}".replace(",", " ").replace(".", ",")


def _style_row(sheet: Worksheet, row: int, *, fill: PatternFill | None = None, bold: bool = False) -> None:
    for column in range(1, len(HEADERS) + 1):
        cell = sheet.cell(row=row, column=column)
        cell.border = BORDER
        if fill is not None:
            cell.fill = fill
        if bold:
            cell.font = BOLD


def _write_purchase(sheet: Worksheet, result: CalculationResult, title: str) -> None:
    sheet["A1"] = f"Закупочный лист: {title}"
    sheet["A1"].font = TITLE
    sheet["A2"] = (
        f"Сформирован {date.today():%d.%m.%Y}; запас {result.margin_percent}% учтён в графе «Нужно с запасом»"
    )
    header_row = 4
    for column, (header, width) in enumerate(zip(HEADERS, WIDTHS, strict=True), start=1):
        cell = sheet.cell(row=header_row, column=column, value=header)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(column)].width = width
    _style_row(sheet, header_row, fill=HEADER_FILL, bold=True)

    row = header_row + 1
    for section_number, (section, lines) in enumerate(
        groupby(result.lines, key=lambda line: line.section), start=1
    ):
        sheet.cell(row=row, column=1, value=f"{section_number}.")
        sheet.cell(row=row, column=2, value=section.title)
        _style_row(sheet, row, fill=SECTION_FILL, bold=True)
        row += 1
        for position, line in enumerate(lines, start=1):
            values = [
                f"{section_number}.{position}",
                line.name,
                line.unit,
                _number(line.required),
                _number(line.product.pack_size) if line.product else None,
                line.packs,
                _number(line.purchase_quantity),
                _number(line.product.pack_price) if line.product else None,
                _number(line.cost),
                line.product.supplier if line.product else "",
            ]
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row=row, column=column, value=value)
                if column in (4, 5, 7):
                    cell.number_format = QUANTITY_FORMAT
                elif column in (8, 9):
                    cell.number_format = MONEY_FORMAT
            _style_row(sheet, row)
            row += 1

    sheet.cell(row=row, column=2, value="Итого")
    total = sheet.cell(row=row, column=9, value=_number(result.total_cost))
    total.number_format = MONEY_FORMAT
    _style_row(sheet, row, bold=True)
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(HEADERS))}{row - 1}"


def _write_breakdown(sheet: Worksheet, result: CalculationResult) -> None:
    headers = ["Блюдо", "Порций", "Продукт", "Ед. изм.", "Брутто на все порции"]
    for column, (header, width) in enumerate(zip(headers, [34, 9, 40, 9, 20], strict=True), start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row, item in enumerate(result.breakdown, start=2):
        sheet.cell(row=row, column=1, value=item.dish)
        sheet.cell(row=row, column=2, value=item.portions)
        sheet.cell(row=row, column=3, value=item.ingredient)
        sheet.cell(row=row, column=4, value=item.unit)
        sheet.cell(row=row, column=5, value=_number(item.quantity)).number_format = QUANTITY_FORMAT
    sheet.freeze_panes = "A2"


def write_report(path: Path, result: CalculationResult, title: str) -> Path:
    workbook = Workbook()
    del workbook[workbook.sheetnames[0]]  # пустой лист по умолчанию — листы создаются явно
    _write_purchase(workbook.create_sheet("Закупка"), result, title)
    _write_breakdown(workbook.create_sheet("По блюдам"), result)
    warnings = workbook.create_sheet("Предупреждения")
    warnings.column_dimensions["A"].width = 120
    for row, text in enumerate(result.warnings or ["Предупреждений нет"], start=1):
        warnings.cell(row=row, column=1, value=text)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
